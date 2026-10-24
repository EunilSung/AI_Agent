"""파일 바이트에서 요약할 글을 뽑는다. Talk 에서 받은 문서를 요약하기 전 단계다.

지원 형식
  글     txt md csv tsv log json xml html htm   UTF-8(BOM 포함) → UTF-16(BOM) → cp949 → 거의 맞는 UTF-8 순서로 읽는다.
                                                html 은 태그를 걷는다
  PDF    pdf                                    pypdf. 암호가 걸렸으면 encrypted_file, 글자가 거의 없으면(스캔) no_text
  Office docx pptx                               python-docx · python-pptx
         xlsx                                    직접 흘려 읽는다(openpyxl 은 날짜 서식 판정·날짜 변환에만 쓴다)
  한글   hwpx hwp(5.0)                           hwpx 는 zip 속 XML, hwp 는 OLE 속 레코드(olefile). 옛 HWP 3.x 는 지원하지 않는다

상한
  요약은 앞부분 조각까지만 읽으므로 글은 MAX_TEXT_CHARS 자까지만 뽑고 멈춘다(모든 형식). 작은 파일이 엄청나게 커지는
  압축 폭탄도 막는다 — HWP 본문은 풀린 크기, zip 은 항목 수·항목 하나·전체 크기, 흘려 읽는 XML 은 읽는 바이트 수에 끝이 있다.
  라이브러리가 통째로 트리로 펼치는 docx·pptx 의 XML 은 이름·확장자가 아니라 첫머리로 가린다(may_be_xml) — 라이브러리는
  [Content_Types].xml 이 정한 형식으로 읽으므로 이름을 바꾼 XML 도 펼친다. xlsx·hwpx 는 트리를 만들지 않고 흘려 읽는다
  (feed_xml) — 들고 있는 것(지금 문단·행, 공유 문자열 등)마다 끝이 있어 XML 모양이 어떻든 메모리가 묶인다.
  상한에 닿아 뒷부분을 읽지 않았으면 Extracted.truncated 로 알린다.

실패는 ExtractError(code, message) 로 알린다. code 는 앱이 이유를 보여 줄 수 있는 SummaryErrorCode 와 같은 값이다.
LLM 을 부르지 않고 chroma 도 쓰지 않는다 — 테스트에서 따로 불러 쓸 수 있다.
"""

import io
import posixpath
import re
import struct
import zipfile
import zlib
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from html.parser import HTMLParser
from xml.etree import ElementTree

UNSUPPORTED_FILE = "unsupported_file"
NO_TEXT = "no_text"
ENCRYPTED_FILE = "encrypted_file"
BROKEN_FILE = "broken_file"

TEXT_EXTENSIONS = {"txt", "md", "markdown", "csv", "tsv", "log", "json", "xml"}
HTML_EXTENSIONS = {"html", "htm"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | HTML_EXTENSIONS | {"pdf", "docx", "xlsx", "pptx", "hwpx", "hwp"}

# 뽑는 글의 끝(모든 형식). 요약은 앞 12조각(약 84,000자)까지만 읽으므로, 공백을 줄인 뒤에도 넉넉히 남는 길이에서 멈춘다.
MAX_TEXT_CHARS = 200_000
# 글 파일이 UTF-8 로도 cp949 로도 딱 맞지 않을 때, 못 읽는 글자(�)가 이 비율 이하면 UTF-8 로 보고 그 바이트만 � 로 바꾼다.
# 여러 출처를 이어 붙인 로그, 한글 중간에서 잘린 파일처럼 몇 바이트만 어긋난 경우다.
MAX_BAD_CHAR_RATIO = 0.01
# html 은 이만큼씩 나눠 파서에 넣고, 모은 글이 MAX_TEXT_CHARS 의 이 배수를 넘으면 그만 넣는다(공백·줄을 줄이기 전 길이라 넉넉히).
HTML_FEED_CHARS = 64 * 1024
HTML_RAW_FACTOR = 2
# 엑셀은 시트마다 이 행 수까지만 읽는다.
MAX_SHEET_ROWS = 2000
MAX_SHEET_COLUMNS = 60
# 엑셀은 모든 시트를 합쳐 빈 행을 포함해 이 행 수까지만 훑는다. openpyxl 은 건너뛴 행 번호를 빈 행으로 채워 내놓으므로,
# 백만 번째 행에 값이 하나 있는 5KB 파일도 빈 행 백만 개를 돈다. 데이터가 그보다 뒤에서 시작하는 시트는 사실상 없다.
MAX_SCAN_ROWS = 50_000
# PDF 는 쪽마다 평균 이 글자 수보다 적으면 스캔한 그림 문서로 본다(쪽 번호·머리말 정도만 글자로 남는다).
MIN_PDF_CHARS_PER_PAGE = 10
MIN_PDF_CHARS = 20
# PDF 는 이 쪽 수까지만 읽는다. 빈 쪽 수만 장짜리 파일에서 글자 상한에 닿지 않고 계속 도는 일을 막는다.
MAX_PDF_PAGES = 2_000
# zip 으로 싼 문서를 풀었을 때 크기의 끝(작은 파일이 엄청나게 커지는 압축 폭탄).
# 라이브러리(python-docx·python-pptx)가 통째로 트리로 펼칠 수 있는 항목 — 이름이 아니라 첫머리가 XML 일 수 있는 항목
# (may_be_xml) — 은 메모리에서 열 배 가까이 불어나므로 하나가 MAX_ZIP_PART_BYTES, 합쳐서 MAX_UNZIPPED_BYTES 까지다.
# 나머지 — 직접 흘려 읽는 XML(xlsx 전부, hwpx 구역)과 XML 이 아닌 항목(그림 등, 읽어도 제 크기만큼만 차지한다) — 는
# 합쳐서 MAX_UNZIPPED_OTHER_BYTES 까지다. 20MB 짜리 엑셀은 시트가 200MB 가까이 풀리기도 한다.
MAX_ZIP_PART_BYTES = 16 * 1024 * 1024
MAX_UNZIPPED_BYTES = 48 * 1024 * 1024
MAX_UNZIPPED_OTHER_BYTES = 256 * 1024 * 1024
# zip 항목 수의 끝. 항목마다 첫머리를 풀어 보므로(may_be_xml_entry) 수만 개짜리 zip 에 오래 붙잡히지 않게 한다.
MAX_ZIP_ENTRIES = 10_000
# zip 중앙 디렉터리(항목 목록) 크기의 끝. zipfile 은 열 때 항목마다 ZipInfo(수백 바이트)를 만들어, 빈 항목 수십만 개를 담은
# 20MB 파일로도 백 MB 넘게 쓴다. 그래서 열기 전에 끝 레코드가 밝힌 크기를 본다(open_zip). 항목 MAX_ZIP_ENTRIES 개에 넉넉하다.
MAX_ZIP_DIRECTORY_BYTES = 2 * 1024 * 1024
# XML 일 수 있는지 볼 때 풀어 보는 항목 첫머리의 바이트 수.
XML_SNIFF_BYTES = 4096
# 흘려 읽는 XML 을 파서에 한 번에 넣는 바이트 수.
XML_FEED_BYTES = 64 * 1024
# xlsx 를 읽는 바이트 수의 끝. 통합 문서·관계·스타일은 항목마다 MAX_ZIP_PART_BYTES, 공유 문자열은
# MAX_XLSX_STRINGS_READ_BYTES, 시트는 모두 합쳐 MAX_XLSX_SHEETS_READ_BYTES 까지 읽고, 넘으면 앞부분만 쓴다.
MAX_XLSX_STRINGS_READ_BYTES = 32 * 1024 * 1024
MAX_XLSX_SHEETS_READ_BYTES = 64 * 1024 * 1024
# 공유 문자열은 칸이 번호로 가리키므로 리스트로 들고 있어야 한다. 문자열 수와 글자 수에 끝을 둔다 — 넘는 문자열을 가리키는
# 칸은 빈칸이 되고 truncated 로 알린다. 고르게 반복되는 XML 은 수백 배로 줄어들어 작은 파일에도 문자열 수백만 개를 실을 수 있다.
MAX_XLSX_STRINGS = 500_000
MAX_XLSX_STRING_CHARS = 4_000_000
# 칸 하나의 글자 수 끝(엑셀의 칸 글자 상한과 같다). 넘는 부분은 버린다.
MAX_CELL_CHARS = 32_767
# 시트 목록·관계·칸 서식·숫자 서식 수와 시트 이름·서식 글자 수의 끝. 넘는 것은 읽지 않는다.
MAX_XLSX_SHEETS = 1_000
MAX_XLSX_RELATIONSHIPS = 10_000
MAX_XLSX_STYLES = 100_000
MAX_SHEET_NAME_CHARS = 100
MAX_FORMAT_CHARS = 255
# hwpx 구역 XML 은 통째로 펼치지 않고 흘려 읽으며, 구역을 합쳐 이 바이트 수까지만 읽는다(넘으면 앞부분만).
MAX_HWPX_READ_BYTES = 32 * 1024 * 1024
# HWP 본문 구역을 풀었을 때 크기 합의 끝. 넘는 부분은 풀지 않는다 — deflate 는 1000 배 가까이 줄어들어 20MB 파일이 수 GB 로 풀린다.
MAX_HWP_BODY_BYTES = 32 * 1024 * 1024

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK\x03\x04"
HWP_SIGNATURE = b"HWP Document File"


class ExtractError(Exception):
    """글을 뽑지 못했다. code 는 SummaryErrorCode 중 하나, message 는 서버 기록용 영어 설명이다."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class Extracted:
    """뽑은 글. truncated 는 글자 수·크기 상한에 닿아 문서 뒷부분을 읽지 않았다는 뜻이다."""

    text: str
    truncated: bool


class TextSink:
    """뽑은 글을 줄마다 모은다. 합친 길이가 limit 자에 닿으면 넘는 부분을 버리고, 그 뒤로는 받지 않는다."""

    def __init__(self, limit: int):
        self.limit = limit
        self.lines: list[str] = []
        self.length = 0
        self.full = False
        self.truncated = False

    def add(self, line: str) -> bool:
        """줄 하나를 더한다. 더 받을 수 있으면 True, 상한에 닿아 그만 읽어야 하면 False."""
        if self.full:
            return False
        separator = 1 if self.lines else 0
        room = self.limit - self.length - separator
        if len(line) > room:
            if room > 0:
                self.lines.append(line[:room])
                self.length = self.limit
            self.full = self.truncated = True
            return False
        self.lines.append(line)
        self.length += separator + len(line)
        return True

    def text(self) -> str:
        return "\n".join(self.lines)


def extension_of(name: str) -> str:
    """경로를 뗀 파일 이름의 확장자(소문자). 없으면 빈 문자열."""
    base = name.replace("\\", "/").rsplit("/", 1)[-1]
    return base.rsplit(".", 1)[-1].lower() if "." in base.strip(".") else ""


def extract_text(name: str, data: bytes) -> str:
    """extract 의 글만 돌려준다."""
    return extract(name, data).text


def extract(name: str, data: bytes) -> Extracted:
    """파일 이름의 확장자로 형식을 가려 글을 뽑는다. 확장자가 없으면 내용 첫머리로 짐작한다.

    이름이 hwp 인데 내용이 zip 이면(hwpx 를 .hwp 로 저장한 경우) hwpx 로 읽는다.
    """
    if not data:
        raise ExtractError(NO_TEXT, "empty file")
    ext = extension_of(name) or sniff(data)
    if ext == "hwp" and data.startswith(ZIP_MAGIC):
        ext = "hwpx"
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractError(UNSUPPORTED_FILE, f"unsupported file type: {ext or 'unknown'}")

    sink = TextSink(MAX_TEXT_CHARS)
    try:
        if ext in TEXT_EXTENSIONS:
            sink.add(decode_text(data))
        elif ext in HTML_EXTENSIONS:
            html_text(decode_text(data), sink)
        elif ext == "pdf":
            pdf_text(data, sink)
        elif ext == "hwp":
            hwp_text(data, sink)
        else:
            if data.startswith(OLE_MAGIC):
                office_in_ole(data)
            check_zip(data, LOADED_WHOLE[ext])
            {"docx": docx_text, "xlsx": xlsx_text, "pptx": pptx_text, "hwpx": hwpx_text}[ext](data, sink)
    except ExtractError:
        raise
    except Exception as e:
        raise ExtractError(BROKEN_FILE, f"cannot read {ext}: {e.__class__.__name__}: {e}") from e

    text = sink.text()
    if not re.search(r"\S", text):
        raise ExtractError(NO_TEXT, "no text in file")
    return Extracted(text, sink.truncated)


def sniff(data: bytes) -> str:
    """확장자가 없는 파일의 형식을 첫머리로 짐작한다. 모르면 빈 문자열."""
    if b"%PDF-" in data[:1024]:
        return "pdf"
    if data.startswith(OLE_MAGIC):
        return "hwp" if HWP_SIGNATURE in data else ""
    if data.startswith(ZIP_MAGIC):
        try:
            names = set(open_zip(data).namelist())
        except zipfile.BadZipFile:
            return ""
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "ppt/presentation.xml" in names:
            return "pptx"
        if any(n.startswith("Contents/section") for n in names):
            return "hwpx"
        return ""
    return "" if b"\x00" in data[:4096] else "txt"


# ---------- 글 파일 ----------

def decode_text(data: bytes) -> str:
    """UTF-8(BOM 포함) → UTF-16(BOM 이 있을 때) → cp949 순서로 읽는다.

    셋 다 딱 맞지 않아도 못 읽는 글자가 MAX_BAD_CHAR_RATIO 이하면 UTF-8 로 읽고 그 바이트만 � 로 바꾼다.
    cp949 를 먼저 보므로, 영문 사이에 한글이 조금 섞인 cp949 파일을 UTF-8 로 잘못 읽지 않는다.
    """
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    text = data.decode("utf-8-sig", errors="replace")
    if text.count("�") <= len(text) * MAX_BAD_CHAR_RATIO:
        return text
    raise ExtractError(BROKEN_FILE, "text is neither UTF-8 nor CP949")


class _HtmlText(HTMLParser):
    """html 의 보이는 글만 모은다. script·style 은 버리고, 덩어리 태그 앞뒤에서 줄을 바꾼다."""

    BLOCKS = {
        "p", "div", "br", "li", "ul", "ol", "tr", "table", "section", "article", "header", "footer",
        "h1", "h2", "h3", "h4", "h5", "h6", "pre", "blockquote", "hr", "dd", "dt", "title",
    }
    HIDDEN = {"script", "style", "noscript", "template"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden = 0
        self.length = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.HIDDEN:
            self.hidden += 1
        if tag in self.BLOCKS:
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag):
        if tag in self.HIDDEN:
            self.hidden = max(0, self.hidden - 1)
        if tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)
            self.length += len(data)


def html_text(html: str, sink: TextSink):
    """html 을 조금씩 파서에 넣어 보이는 글을 줄마다 sink 에 담는다. 모은 글이 넉넉히 쌓이면 남은 html 은 읽지 않는다."""
    parser = _HtmlText()
    for start in range(0, len(html), HTML_FEED_CHARS):
        if parser.length > sink.limit * HTML_RAW_FACTOR:
            sink.truncated = True
            break
        parser.feed(html[start:start + HTML_FEED_CHARS])
    else:
        parser.close()
    for line in "".join(parser.parts).splitlines():
        line = re.sub(r"[ \t ]+", " ", line).strip()
        if not sink.add(re.sub(r"^\|\s*", "", line)):
            return


# ---------- PDF ----------

def pdf_text(data: bytes, sink: TextSink):
    """쪽마다 글을 뽑아 빈 줄을 사이에 두고 담는다. 글자 상한이나 MAX_PDF_PAGES 에 닿으면 남은 쪽은 읽지 않는다."""
    from pypdf import PdfReader
    from pypdf.errors import DependencyError, FileNotDecryptedError

    if b"%PDF-" not in data[:1024]:
        raise ExtractError(BROKEN_FILE, "not a PDF file")
    pages = visible = 0
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractError(ENCRYPTED_FILE, "PDF is password protected")
        for page in reader.pages:
            if pages >= MAX_PDF_PAGES:
                sink.truncated = True
                break
            text = (page.extract_text() or "").strip()
            visible += len(re.sub(r"\s", "", text))
            pages += 1
            if (pages > 1 and not sink.add("")) or not sink.add(text):
                break
    except ExtractError:
        raise
    except (FileNotDecryptedError, DependencyError) as e:
        raise ExtractError(ENCRYPTED_FILE, f"PDF is encrypted: {e}") from e
    except Exception as e:
        raise ExtractError(BROKEN_FILE, f"cannot read PDF: {e.__class__.__name__}: {e}") from e

    if visible < max(MIN_PDF_CHARS, MIN_PDF_CHARS_PER_PAGE * pages):
        raise ExtractError(NO_TEXT, f"PDF has almost no text ({visible} chars in {pages} pages)")


# ---------- Office (docx · xlsx · pptx) ----------

def office_in_ole(data: bytes):
    """docx·xlsx·pptx 인데 OLE 로 싸여 있으면 암호 걸린 Office 문서이거나 이름만 바꾼 옛 형식이다."""
    import olefile

    try:
        ole = olefile.OleFileIO(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(BROKEN_FILE, f"broken OLE container: {e}") from e
    with ole:
        if ole.exists("EncryptionInfo") or ole.exists("EncryptedPackage"):
            raise ExtractError(ENCRYPTED_FILE, "Office document is password protected")
    raise ExtractError(UNSUPPORTED_FILE, "legacy binary Office document")


# BOM 과 그 뒤를 읽는 인코딩. UTF-32 를 UTF-16 보다 먼저 본다(UTF-32LE 의 BOM 이 UTF-16LE 의 BOM 으로 시작한다).
XML_BOMS = (
    (b"\x00\x00\xfe\xff", "utf-32-be"),
    (b"\xff\xfe\x00\x00", "utf-32-le"),
    (b"\xef\xbb\xbf", "utf-8"),
    (b"\xfe\xff", "utf-16-be"),
    (b"\xff\xfe", "utf-16-le"),
)
# BOM 없이 libxml2(lxml)가 첫 네 바이트로 알아보는 UTF-32·UTF-16·EBCDIC 의 '<'·'<?'.
XML_UNMARKED_STARTS = (
    b"\x00\x00\x00<", b"<\x00\x00\x00", b"\x00\x00<\x00", b"\x00<\x00\x00",
    b"\x00<\x00?", b"<\x00?\x00", b"\x4c\x6f\xa7\x94",
)


def may_be_xml(head: bytes) -> bool:
    """항목 첫머리(head)가 XML 문서의 시작일 수 있는지. 아니라고 할 수 있을 때만 False 다.

    XML 은 BOM 과 공백(빈칸·탭·줄바꿈) 뒤에 '<' 로 시작해야 한다. lxml 이 BOM 없이 알아보는 UTF-16·UTF-32·EBCDIC 의
    첫머리도 XML 로 본다. head 가 모두 공백이면 뒤에 무엇이 올지 모르므로 XML 로 본다. 그림(PNG·JPEG 등)·동영상·OLE·zip 은
    첫 바이트가 '<' 가 아니라 False 다.
    """
    if head.startswith(XML_UNMARKED_STARTS):
        return True
    for bom, encoding in XML_BOMS:
        if head.startswith(bom):
            text = head[len(bom):].decode(encoding, errors="ignore")
            break
    else:
        text = head.decode("latin-1")
    rest = text.lstrip(" \t\r\n")
    return not rest or rest.startswith("<")


def may_be_xml_entry(archive: zipfile.ZipFile, info: zipfile.ZipInfo) -> bool:
    """zip 항목이 XML 일 수 있는지 첫 XML_SNIFF_BYTES 바이트를 풀어 본다. 이름·확장자는 보지 않는다.

    python-docx·python-pptx 는 [Content_Types].xml 이 정한 형식으로 항목을 읽으므로, word/document.xml 을 .bin 으로
    바꾸고 형식만 밝혀 두면 그 항목도 통째로 트리로 펼친다. 이름으로 가리면 그런 항목이 큰 예산으로 빠져나간다.
    """
    if info.is_dir():
        return False
    with archive.open(info) as stream:
        return may_be_xml(stream.read(XML_SNIFF_BYTES))


# 형식마다, 라이브러리가 통째로 메모리에 트리로 펼칠 수 있는 zip 항목인지. python-docx·python-pptx 는 XML 항목을 모두
# 트리로 펼친다. xlsx·hwpx 는 직접 흘려 읽는다 — hwpx 의 manifest 는 바이트 그대로 읽어 들이므로 한 항목 상한만 건다.
LOADED_WHOLE: dict[str, Callable[[zipfile.ZipFile, zipfile.ZipInfo], bool]] = {
    "docx": may_be_xml_entry,
    "pptx": may_be_xml_entry,
    "xlsx": lambda archive, info: False,
    "hwpx": lambda archive, info: info.filename == "META-INF/manifest.xml",
}


ZIP_END = b"PK\x05\x06"
ZIP64_END_LOCATOR = b"PK\x06\x07"
ZIP64_END = b"PK\x06\x06"


def zip_directory_size(data: bytes) -> int | None:
    """zip 끝 레코드가 밝힌 중앙 디렉터리 크기(바이트). 끝 레코드를 찾지 못하면 None.

    zipfile(_EndRecData) 이 찾는 방법과 같다 — 파일 끝 22바이트가 설명 없는 끝 레코드면 그것을, 아니면 마지막 64KB(설명 칸)
    안의 마지막 끝 레코드를 쓴다. 바로 앞에 ZIP64 위치 레코드가 있으면 그 레코드가 가리키는 곳(없으면 바로 앞)의 ZIP64 끝
    레코드 값을 쓴다. zipfile 은 이 크기만큼 중앙 디렉터리를 읽어 항목마다 ZipInfo 를 만든다. ZIP64 끝 레코드를 찾지 못하면
    32비트 값(ZIP64 면 0xFFFFFFFF)이 남아 상한에 걸린다.
    """
    end = len(data) - 22
    if not (end >= 0 and data[end:end + 4] == ZIP_END and data[-2:] == b"\x00\x00"):
        end = data.rfind(ZIP_END, max(0, len(data) - 22 - 65535))
        if end < 0 or end + 22 > len(data):
            return None
    (size,) = struct.unpack_from("<I", data, end + 12)
    locator = end - 20
    if locator < 0 or data[locator:locator + 4] != ZIP64_END_LOCATOR:
        return size
    (relative,) = struct.unpack_from("<Q", data, locator + 8)
    for record in (relative, locator - 56):
        if 0 <= record <= len(data) - 56 and data[record:record + 4] == ZIP64_END:
            (size,) = struct.unpack_from("<Q", data, record + 40)
            break
    return size


def open_zip(data: bytes) -> zipfile.ZipFile:
    """data 를 zip 으로 연다. 중앙 디렉터리가 MAX_ZIP_DIRECTORY_BYTES 보다 크면 열지 않고 broken_file 이다."""
    size = zip_directory_size(data)
    if size is not None and size > MAX_ZIP_DIRECTORY_BYTES:
        raise ExtractError(BROKEN_FILE, f"zip directory is too large ({size} bytes)")
    return zipfile.ZipFile(io.BytesIO(data))


def check_zip(data: bytes, loaded_whole: Callable[[zipfile.ZipFile, zipfile.ZipInfo], bool]):
    """zip 으로 싼 문서인지, 풀었을 때 지나치게 커지지 않는지 본다(작은 파일이 엄청나게 커지는 압축 폭탄).

    항목은 MAX_ZIP_ENTRIES 개까지. loaded_whole(archive, 항목) 이 참인 항목은 하나가 MAX_ZIP_PART_BYTES, 합쳐서
    MAX_UNZIPPED_BYTES 까지, 나머지는 합쳐서 MAX_UNZIPPED_OTHER_BYTES 까지. 크기는 zip 이 밝힌 값이고, zipfile 은
    그보다 많이 풀어내지 않는다.
    """
    try:
        archive = open_zip(data)
    except zipfile.BadZipFile as e:
        raise ExtractError(BROKEN_FILE, f"not a zip package: {e}") from e
    infos = archive.infolist()
    if len(infos) > MAX_ZIP_ENTRIES:
        raise ExtractError(BROKEN_FILE, f"zip package has too many entries ({len(infos)})")
    tree = other = 0
    for info in infos:
        if not loaded_whole(archive, info):
            other += info.file_size
            continue
        if info.file_size > MAX_ZIP_PART_BYTES:
            raise ExtractError(BROKEN_FILE, f"zip entry {info.filename} is too large when unpacked")
        tree += info.file_size
    if tree > MAX_UNZIPPED_BYTES or other > MAX_UNZIPPED_OTHER_BYTES:
        raise ExtractError(BROKEN_FILE, "zip package is too large when unpacked")


def row_line(cells) -> str:
    values = [c.strip() for c in cells]
    while values and not values[-1]:
        values.pop()
    return " | ".join(values)


def docx_text(data: bytes, sink: TextSink):
    """본문 문단과 표를 문서 순서대로. 표는 행마다 한 줄, 칸은 ` | ` 로 잇는다(합친 칸은 한 번만)."""
    from docx import Document
    from docx.table import Table

    document = Document(io.BytesIO(data))
    for block in document.iter_inner_content():
        lines = ["", *table_lines(block), ""] if isinstance(block, Table) else [block.text]
        for line in lines:
            if not sink.add(line):
                return


def table_lines(table) -> Iterator[str]:
    for row in table.rows:
        seen, cells = set(), []
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(cell.text.replace("\n", " "))
        line = row_line(cells)
        if line:
            yield line


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        text = value.isoformat()
        return text[:-9] if text.endswith("T00:00:00") else text
    return str(value).replace("\n", " ")


def xlsx_text(data: bytes, sink: TextSink):
    """시트마다 `## 시트: 이름` 아래 행을 ` | ` 로 잇는다. 값은 수식이 아니라 마지막으로 계산된 값을 쓴다.

    값이 있는 행이 없는 시트는 건너뛴다. 시트마다 MAX_SHEET_ROWS 행, 모든 시트를 합쳐 MAX_SCAN_ROWS 행(빈 행 포함)까지 본다.

    openpyxl 로 읽지 않는다. openpyxl 은 read_only 에서도 공유 문자열·통합 문서·차트 시트를 통째로 펼치고, 흘려 읽는
    시트에서도 행 밖의 요소를 버리지 않아(같은 번호가 되풀이되는 행도) 수십 KB 파일로 수 GB 를 쓰게 할 수 있다. 어느 항목을
    통째로 펼칠지도 [Content_Types].xml 로 정해 이름으로 미리 가릴 수 없다. 여기서는 모든 XML 을 트리 없이 흘려 읽고
    (feed_xml), 들고 있는 것(시트 목록·칸 서식·공유 문자열·지금 행)과 읽는 바이트 수마다 끝을 둔다. 어느 항목을 읽을지는
    [Content_Types].xml 이 아니라 관계(.rels)로 정한다.
    """
    archive = zipfile.ZipFile(io.BytesIO(data))
    book = read_xlsx_book(archive, sink)
    budget = _ByteBudget(MAX_XLSX_SHEETS_READ_BYTES)
    scan_left = MAX_SCAN_ROWS
    first_sheet = True
    for title, part in book.sheets:
        if scan_left <= 0:
            break
        if budget.left <= 0:
            sink.truncated = True
            break
        rows = _SheetRows(title, book, scan_left, sink, first_sheet)
        if not read_part(archive, part, rows, budget):
            sink.truncated = True
        if rows.stopped:
            return
        if rows.rows:
            first_sheet = False
        scan_left -= rows.scanned


XLSX_WORKBOOK = "xl/workbook.xml"
XLSX_SHARED_STRINGS = "xl/sharedStrings.xml"
XLSX_STYLES = "xl/styles.xml"
XLSX_CELL_REFERENCE = re.compile(r"\$?([A-Za-z]{1,3})\$?\d+")


def to_int(value) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def rels_path_of(part: str) -> str:
    """zip 항목 part 의 관계 파일 이름(`폴더/_rels/이름.rels`). part 가 빈 문자열이면 꾸러미 전체의 `_rels/.rels`."""
    folder, base = posixpath.split(part)
    return posixpath.join(folder, "_rels", f"{base}.rels")


class _Relationships:
    """관계 파일(.rels)을 흘려 읽어 관계 id → (종류, zip 속 이름)을 모은다. 바깥 링크(TargetMode=External)는 뺀다.

    대상 경로는 openpyxl·python-docx 와 같이 맞춘다 — '/' 로 시작하면 꾸러미 맨 위에서, 아니면 `_rels` 의 윗 폴더에서 찾는다.
    """

    def __init__(self, rels_path: str):
        self.base = posixpath.split(posixpath.dirname(rels_path))[0]
        self.items: dict[str, tuple[str, str]] = {}

    def start(self, tag, attrib):
        if local_name(tag) != "Relationship" or len(self.items) >= MAX_XLSX_RELATIONSHIPS:
            return
        if attrib.get("TargetMode") == "External":
            return
        target = attrib.get("Target", "")
        name = target[1:] if target.startswith("/") else posixpath.normpath(posixpath.join(self.base, target))
        self.items.setdefault(attrib.get("Id", ""), (attrib.get("Type", ""), name))

    def end(self, tag):
        pass

    def data(self, text):
        pass


class _Workbook:
    """통합 문서 XML 에서 시트 목록((이름, 관계 id), MAX_XLSX_SHEETS 개까지)과 1904 날짜 체계인지를 읽는다."""

    def __init__(self):
        self.sheets: list[tuple[str, str]] = []
        self.date1904 = False

    def start(self, tag, attrib):
        name = local_name(tag)
        if name == "sheet" and len(self.sheets) < MAX_XLSX_SHEETS:
            rid = next((value for key, value in attrib.items() if key.startswith("{") and local_name(key) == "id"), "")
            if rid:
                self.sheets.append((attrib.get("name", "")[:MAX_SHEET_NAME_CHARS], rid))
        elif name == "workbookPr":
            self.date1904 = attrib.get("date1904", "").strip().lower() in ("1", "true")

    def end(self, tag):
        pass

    def data(self, text):
        pass


class _Styles:
    """스타일 XML 에서 숫자 서식(numFmts)과 칸 서식(cellXfs)마다의 숫자 서식 번호를 읽는다(각각 MAX_XLSX_STYLES 개까지)."""

    def __init__(self):
        self.depth = 0
        self.section = ""
        self.custom: dict[int, str] = {}
        self.cell_formats: list[int] = []

    def start(self, tag, attrib):
        self.depth += 1
        name = local_name(tag)
        if self.depth == 2:
            self.section = name
        elif self.depth == 3 and self.section == "numFmts" and name == "numFmt" and len(self.custom) < MAX_XLSX_STYLES:
            number = to_int(attrib.get("numFmtId"))
            if number is not None:
                self.custom[number] = attrib.get("formatCode", "")[:MAX_FORMAT_CHARS]
        elif self.depth == 3 and self.section == "cellXfs" and name == "xf" and len(self.cell_formats) < MAX_XLSX_STYLES:
            self.cell_formats.append(to_int(attrib.get("numFmtId", 0)) or 0)

    def end(self, tag):
        self.depth -= 1

    def data(self, text):
        pass

    def date_styles(self) -> tuple[set[int], set[int]]:
        """(날짜·시간 서식인 칸 서식 번호들, 그중 경과 시간([h]:mm 등) 서식인 번호들). openpyxl 이 값을 바꾸는 기준과 같다."""
        from openpyxl.styles.numbers import builtin_format_code, is_date_format, is_timedelta_format

        dates, durations = set(), set()
        for index, number in enumerate(self.cell_formats):
            code = self.custom[number] if number in self.custom else builtin_format_code(number)
            if is_date_format(code):
                dates.add(index)
            if is_timedelta_format(code):
                durations.add(index)
        return dates, durations


class _SharedStrings:
    """공유 문자열을 흘려 읽는다. MAX_XLSX_STRINGS 개, 합쳐서 MAX_XLSX_STRING_CHARS 자까지만 모으고 넘으면 full 이다.

    문자열마다 `<t>` 의 글(서식 조각 `<r>` 속 `<t>` 포함)을 잇고, 읽는 법 표시(`<rPh>`)는 뺀다 — openpyxl 과 같다.
    """

    def __init__(self):
        self.strings: list[str] = []
        self.chars = 0
        self.full = False
        self.parts: list[str] | None = None
        self.phonetic = 0
        self.in_text = False

    def start(self, tag, attrib):
        name = local_name(tag)
        if name == "si":
            if len(self.strings) >= MAX_XLSX_STRINGS:
                self.full = True
                raise _StopReading
            self.parts = []
        elif name == "rPh":
            self.phonetic += 1
        elif name == "t" and self.parts is not None and not self.phonetic:
            self.in_text = True

    def data(self, text):
        if not self.in_text:
            return
        if self.chars + len(text) > MAX_XLSX_STRING_CHARS:
            self.full = True
            raise _StopReading
        self.parts.append(text)
        self.chars += len(text)

    def end(self, tag):
        name = local_name(tag)
        if name == "t":
            self.in_text = False
        elif name == "rPh":
            self.phonetic = max(0, self.phonetic - 1)
        elif name == "si" and self.parts is not None:
            self.strings.append("".join(self.parts).replace("x005F_", ""))
            self.parts = None


class _Book:
    """xlsx 에서 시트를 읽기 전에 모은 것. sheets 는 (시트 이름, zip 속 시트 항목)이고 차트 시트는 뺐다."""

    def __init__(self, sheets: list[tuple[str, str]], strings: list[str], styles: _Styles, date1904: bool):
        from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, from_excel, from_ISO8601

        self.sheets = sheets
        self.strings = strings
        self.dates, self.durations = styles.date_styles()
        self.epoch = CALENDAR_MAC_1904 if date1904 else CALENDAR_WINDOWS_1900
        self.from_excel = from_excel
        self.from_iso = from_ISO8601

    def value(self, kind: str, style: int | None, raw: str | None, inline: str | None):
        """칸 하나의 값. openpyxl(read_only, data_only)이 내놓는 값과 같게 만든다 — 공유 문자열은 글로, 날짜 서식의 숫자는
        날짜·시간으로 바꾸고, 수식은 마지막으로 계산된 값(`<v>`)을 쓴다. 읽을 수 없는 값은 글 그대로 둔다(openpyxl 은 실패한다).
        """
        if kind == "inlineStr":
            return inline
        if not raw:
            return None
        try:
            if kind == "s":
                index = int(raw)
                return self.strings[index] if 0 <= index < len(self.strings) else ""
            if kind == "b":
                return bool(int(raw))
            if kind == "d":
                return self.from_iso(raw)
            if kind != "n":
                return raw
            number = float(raw) if any(mark in raw for mark in ".Ee") else int(raw)
        except ValueError:
            return raw
        if style not in self.dates:
            return number
        try:
            return self.from_excel(number, self.epoch, timedelta=style in self.durations)
        except (OverflowError, ValueError):
            return "#VALUE!"


def xlsx_relationships(archive: zipfile.ZipFile, names: set[str], part: str) -> dict[str, tuple[str, str]]:
    """part 의 관계(관계 id → (종류, zip 속 이름)). part 가 빈 문자열이면 꾸러미 전체의 관계. 관계 파일이 없으면 빈 dict."""
    path = rels_path_of(part)
    if path not in names:
        return {}
    handler = _Relationships(path)
    read_part(archive, path, handler, _ByteBudget(MAX_ZIP_PART_BYTES))
    return handler.items


def related_part(rels: dict[str, tuple[str, str]], kind: str) -> str | None:
    """관계 종류가 `/kind` 로 끝나는 첫 관계의 zip 속 이름."""
    return next((name for type_, name in rels.values() if type_.endswith("/" + kind)), None)


def read_xlsx_book(archive: zipfile.ZipFile, sink: TextSink) -> _Book:
    """통합 문서(꾸러미 관계의 officeDocument, 없으면 xl/workbook.xml)·그 관계·스타일·공유 문자열을 읽는다.

    통합 문서나 공유 문자열을 끝까지 읽지 못했으면 sink.truncated 를 세운다.
    """
    names = set(archive.namelist())
    workbook_part = related_part(xlsx_relationships(archive, names, ""), "officeDocument") or XLSX_WORKBOOK
    if workbook_part not in names:
        raise ExtractError(BROKEN_FILE, "xlsx has no workbook")
    workbook = _Workbook()
    if not read_part(archive, workbook_part, workbook, _ByteBudget(MAX_ZIP_PART_BYTES)):
        sink.truncated = True
    rels = xlsx_relationships(archive, names, workbook_part)

    styles = _Styles()
    styles_part = related_part(rels, "styles") or XLSX_STYLES
    if styles_part in names:
        read_part(archive, styles_part, styles, _ByteBudget(MAX_ZIP_PART_BYTES))

    strings = _SharedStrings()
    strings_part = related_part(rels, "sharedStrings") or XLSX_SHARED_STRINGS
    if strings_part in names:
        complete = read_part(archive, strings_part, strings, _ByteBudget(MAX_XLSX_STRINGS_READ_BYTES))
        if not complete or strings.full:
            sink.truncated = True

    sheets = []
    for title, rid in workbook.sheets:
        kind, part = rels.get(rid, ("", ""))
        if part in names and "chartsheet" not in kind:
            sheets.append((title, part))
    return _Book(sheets, strings.strings, styles, workbook.date1904)


def column_number(letters: str) -> int:
    number = 0
    for letter in letters.upper():
        number = number * 26 + ord(letter) - ord("A") + 1
    return number


class _SheetRows:
    """시트 XML 하나를 흘려 읽어 값이 있는 행마다 sink 에 한 줄(` | ` 로 이은 칸)로 담는다.

    트리를 만들지 않는다 — 들고 있는 것은 지금 행의 앞 MAX_SHEET_COLUMNS 칸(칸마다 MAX_CELL_CHARS 자)뿐이다.
    행 번호는 r 속성(없거나 이상하면 앞 행 + 1)이다. openpyxl 의 read_only 와 같이, 앞서 내놓은 행 번호보다 작거나 같은
    행은 건너뛰고, 빠진 번호도 빈 행으로 센다(scanned). scan_limit 번째를 넘는 행이 나오면 그만 읽는다.
    처음으로 값이 있는 행 앞에 `## 시트: 이름` 을 적는다(첫 시트가 아니면 빈 줄을 먼저). 시트마다 MAX_SHEET_ROWS 행까지 적고,
    sink 가 차면 stopped 를 세우고 그만 읽는다.
    """

    def __init__(self, title: str, book: _Book, scan_limit: int, sink: TextSink, first_sheet: bool):
        self.title = title
        self.book = book
        self.scan_limit = scan_limit
        self.sink = sink
        self.first_sheet = first_sheet
        self.scanned = 0
        self.rows = 0
        self.stopped = False
        self.next_row = 1
        self.row_number = 0
        self.depth = 0
        self.row_depth = 0
        self.cells: dict[int, object] = {}
        self.column = 0
        self.cell_depth = 0
        self.kind = "n"
        self.style: int | None = 0
        self.value: list[str] | None = None
        self.inline: list[str] | None = None
        self.inline_depth = 0
        self.target: list[str] | None = None
        self.cell_chars = 0
        self.phonetic = 0

    def start(self, tag, attrib):
        self.depth += 1
        name = local_name(tag)
        if self.cell_depth:
            if name == "v" and self.depth == self.cell_depth + 1:
                self.value = self.target = []
            elif name == "is" and self.depth == self.cell_depth + 1:
                self.inline = []
                self.inline_depth = self.depth
            elif name == "rPh":
                self.phonetic += 1
            elif name == "t" and self.inline_depth and not self.phonetic:
                self.target = self.inline
        elif self.row_depth:
            if name == "c" and self.depth == self.row_depth + 1:
                reference = XLSX_CELL_REFERENCE.fullmatch(attrib.get("r", ""))
                self.column = column_number(reference.group(1)) if reference else self.column + 1
                self.cell_depth = self.depth
                self.kind = attrib.get("t", "n")
                self.style = to_int(attrib.get("s", 0))
                self.value = self.inline = self.target = None
                self.inline_depth = self.cell_chars = self.phonetic = 0
        elif name == "row":
            number = to_int(attrib.get("r"))
            self.row_number = number if number is not None else self.row_number + 1
            self.row_depth = self.depth
            self.cells = {}
            self.column = 0

    def data(self, text):
        if self.target is None or self.cell_chars >= MAX_CELL_CHARS:
            return
        text = text[:MAX_CELL_CHARS - self.cell_chars]
        self.target.append(text)
        self.cell_chars += len(text)

    def end(self, tag):
        depth = self.depth
        self.depth -= 1
        if self.cell_depth:
            name = local_name(tag)
            if depth == self.cell_depth:
                self.cell_depth = 0
                self.finish_cell()
            elif depth == self.inline_depth:
                self.inline_depth = 0
                self.target = None
            elif name in ("v", "t"):
                self.target = None
            elif name == "rPh":
                self.phonetic = max(0, self.phonetic - 1)
        elif self.row_depth and depth == self.row_depth:
            self.row_depth = 0
            self.finish_row()

    def finish_cell(self):
        self.target = None
        if not 1 <= self.column <= MAX_SHEET_COLUMNS:
            return
        raw = "".join(self.value) if self.value is not None else None
        inline = "".join(self.inline) if self.inline is not None else None
        self.cells[self.column] = self.book.value(self.kind, self.style, raw, inline)

    def finish_row(self):
        number = self.row_number
        if number > self.scan_limit:
            self.scanned = self.scan_limit
            raise _StopReading
        if number < self.next_row:
            return
        self.next_row = number + 1
        self.scanned = number
        values = [self.cells.get(column) for column in range(1, max(self.cells, default=0) + 1)]
        if all(value is None for value in values):
            return
        line = row_line(cell_text(value) for value in values)
        if not line:
            return
        if not self.rows:
            if not self.first_sheet and not self.sink.add(""):
                self.stop()
            if not self.sink.add(f"## 시트: {self.title}"):
                self.stop()
        self.rows += 1
        if self.rows > MAX_SHEET_ROWS:
            self.sink.truncated = True
            if not self.sink.add(f"… ({MAX_SHEET_ROWS}행 이후 생략)"):
                self.stopped = True
            raise _StopReading
        if not self.sink.add(line):
            self.stop()

    def stop(self):
        self.stopped = True
        raise _StopReading


def pptx_text(data: bytes, sink: TextSink):
    """슬라이드마다 `## 슬라이드 N` 아래 도형 글·표·발표자 노트를 적는다."""
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    for number, slide in enumerate(presentation.slides, 1):
        lines = [f"## 슬라이드 {number}"] if number == 1 else ["", f"## 슬라이드 {number}"]
        for shape in slide.shapes:
            lines.extend(shape_lines(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"발표자 노트: {notes.text.strip()}")
        for line in lines:
            if not sink.add(line):
                return


def shape_lines(shape) -> list[str]:
    from pptx.shapes.group import GroupShape

    if isinstance(shape, GroupShape):
        return [line for child in shape.shapes for line in shape_lines(child)]
    if getattr(shape, "has_table", False) and shape.has_table:
        lines = []
        for row in shape.table.rows:
            line = row_line(cell.text.replace("\n", " ") for cell in row.cells)
            if line:
                lines.append(line)
        return lines
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return [p.strip() for p in shape.text_frame.text.splitlines() if p.strip()]
    return []


# ---------- 한글 (hwpx · hwp) ----------

def local_name(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


class _ByteBudget:
    """여러 스트림에 걸쳐 읽을 수 있는 바이트 수. cut 은 남은 내용이 있는데 끝에 닿아 읽지 않았다는 뜻이다."""

    def __init__(self, limit: int):
        self.left = limit
        self.cut = False


class _LimitedReader:
    """stream 을 budget 이 남은 만큼만 읽는다. 다 쓰면 끝(빈 바이트)을 돌려준다."""

    def __init__(self, stream, budget: _ByteBudget):
        self.stream = stream
        self.budget = budget

    def read(self, size: int = -1) -> bytes:
        if self.budget.left <= 0:
            if self.stream.read(1):
                self.budget.cut = True
            return b""
        size = self.budget.left if size is None or size < 0 else min(size, self.budget.left)
        chunk = self.stream.read(size)
        self.budget.left -= len(chunk)
        return chunk


class _StopReading(Exception):
    """흘려 읽던 XML 을 그만 읽는다(글·행 상한에 닿음). 처리기가 던지고 feed_xml 이 받는다."""


def feed_xml(stream, handler) -> bool:
    """stream 의 XML 을 XML_FEED_BYTES 씩 읽어 handler(start·end·data 를 받는 처리기)에 넘긴다. 트리를 만들지 않는다.

    메모리에 남는 것은 handler 가 모으는 것과, 파서가 아직 다 읽지 못한 토큰 하나(태그·속성)뿐이다 — 글은 조각으로 온다.
    handler 가 _StopReading 을 던져 멈췄으면 False, 끝까지 읽었으면 True. 깨진 XML 이면 ElementTree.ParseError 다.
    """
    parser = ElementTree.XMLParser(target=handler)
    try:
        while chunk := stream.read(XML_FEED_BYTES):
            parser.feed(chunk)
        parser.close()
    except _StopReading:
        return False
    return True


def read_part(archive: zipfile.ZipFile, name: str, handler, budget: _ByteBudget) -> bool:
    """zip 항목 name 을 budget 이 남은 만큼만 handler 에 흘려 넣는다. budget 에 닿아 뒷부분을 읽지 못했으면 False.

    budget 에 닿아 XML 이 중간에 끊긴 것은 깨진 파일로 보지 않는다. handler 가 스스로 멈춘 것은 handler 가 안다.
    """
    with archive.open(name) as stream:
        try:
            feed_xml(_LimitedReader(stream, budget), handler)
        except ElementTree.ParseError:
            if not budget.cut:
                raise
    return not budget.cut


def hwpx_text(data: bytes, sink: TextSink):
    """`Contents/section*.xml` 을 번호 순서대로 흘려 읽어 문단(`<hp:p>`)마다 한 줄로. 표 칸 속 문단도 각자 한 줄이 된다.

    구역 XML 로 트리를 만들지 않는다(hwpx_section_lines). 구역을 합쳐 MAX_HWPX_READ_BYTES 까지만 읽는다.
    """
    archive = zipfile.ZipFile(io.BytesIO(data))
    names = archive.namelist()
    if "META-INF/manifest.xml" in names and b"encryption-data" in archive.read("META-INF/manifest.xml"):
        raise ExtractError(ENCRYPTED_FILE, "HWPX document is password protected")
    sections = sorted(
        (n for n in names if re.fullmatch(r"Contents/section\d+\.xml", n)),
        key=lambda n: int(re.search(r"\d+", n.rsplit("/", 1)[-1]).group()),
    )
    if not sections:
        raise ExtractError(BROKEN_FILE, "HWPX has no sections")

    budget = _ByteBudget(MAX_HWPX_READ_BYTES)
    for name in sections:
        if budget.left <= 0:
            sink.truncated = True
            return
        with archive.open(name) as stream:
            reader = _LimitedReader(stream, budget)
            try:
                complete = hwpx_section_lines(reader, sink)
            except ElementTree.ParseError:
                if not budget.cut:
                    raise
                complete = False
        if budget.cut:
            sink.truncated = True
        if not complete:
            return


def hwpx_section_lines(source, sink: TextSink) -> bool:
    """구역 XML 하나를 흘려 읽어 문단마다 sink 에 담는다. sink 가 차서 멈췄으면 False."""
    paragraphs = _HwpxParagraphs(sink)
    return feed_xml(source, paragraphs) and paragraphs.flush()


# `<hp:t>` 바로 안의 표시 요소를 글로 바꾼 것
HWPX_TEXT_MARKS = {"tab": "\t", "lineBreak": "\n", "fwSpace": " ", "nbSpace": " "}


class _HwpxParagraphs:
    """hwpx 구역 XML 의 문단(`<hp:p>`)마다 글을 모아 sink 에 한 줄로 담는다. 트리를 만들지 않는다.

    글은 `<hp:t>` 바로 안의 글만 모으고, `<hp:t>` 바로 안의 탭·줄바꿈·빈칸 표시는 HWPX_TEXT_MARKS 로 바꾼다.
    문단 안에서 문단이 열리면(표 칸) 그때까지 모은 바깥 문단의 글을 먼저 한 줄로 내보낸다. 한 문단의 글은 sink 에 들어갈
    만큼까지만 모은다. sink 가 차면 _StopReading 을 던진다.
    """

    def __init__(self, sink: TextSink):
        self.sink = sink
        self.current: list[str] = []
        self.length = 0
        self.dropped = False
        self.depth = 0
        self.text_depth = 0

    def start(self, tag, attrib):
        self.depth += 1
        name = local_name(tag)
        if self.text_depth:
            if self.depth == self.text_depth + 1 and name in HWPX_TEXT_MARKS:
                self.append(HWPX_TEXT_MARKS[name])
        elif name == "p":
            if not self.flush():
                raise _StopReading
        elif name == "t":
            self.text_depth = self.depth

    def data(self, text):
        if self.text_depth and self.depth == self.text_depth:
            self.append(text)

    def end(self, tag):
        depth = self.depth
        self.depth -= 1
        if self.text_depth:
            if depth == self.text_depth:
                self.text_depth = 0
        elif local_name(tag) == "p" and not self.flush():
            raise _StopReading

    def append(self, text: str):
        if self.length >= self.sink.limit:
            self.dropped = True
            return
        self.current.append(text)
        self.length += len(text)

    def flush(self) -> bool:
        """모은 문단 글을 한 줄로 내보낸다. sink 가 차서 더 받지 않으면 False."""
        line = "".join(self.current).strip()
        if self.dropped:
            self.sink.truncated = True
        self.current.clear()
        self.length = 0
        self.dropped = False
        return self.sink.add(line) if line else True


HWPTAG_PARA_TEXT = 67
HWP_COMPRESSED = 0x1
HWP_PASSWORD = 0x2
HWP_DISTRIBUTION = 0x4


def hwp_text(data: bytes, sink: TextSink):
    """HWP 5.0: `FileHeader` 의 속성을 보고 `BodyText/Section{n}` 의 문단 글(HWPTAG_PARA_TEXT)을 순서대로 모은다.

    구역을 풀었을 때 크기는 합쳐서 MAX_HWP_BODY_BYTES 까지만 푼다(압축 폭탄). OLE 가 아닌 옛 HWP 3.x 는 지원하지 않는다.
    """
    import olefile

    if not data.startswith(OLE_MAGIC):
        if data.startswith(HWP_SIGNATURE):
            raise ExtractError(UNSUPPORTED_FILE, "HWP 3.x or older binary document")
        raise ExtractError(BROKEN_FILE, "not an HWP 5.0 (OLE) file")
    try:
        ole = olefile.OleFileIO(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(BROKEN_FILE, f"broken HWP container: {e}") from e

    with ole:
        if not ole.exists("FileHeader"):
            raise ExtractError(BROKEN_FILE, "HWP has no FileHeader")
        header = ole.openstream("FileHeader").read()
        if len(header) < 40 or not header.startswith(HWP_SIGNATURE):
            raise ExtractError(BROKEN_FILE, "bad HWP FileHeader")
        (flags,) = struct.unpack_from("<I", header, 36)
        if flags & (HWP_PASSWORD | HWP_DISTRIBUTION):
            raise ExtractError(ENCRYPTED_FILE, "HWP is password protected or distribution-only")

        sections = []
        for entry in ole.listdir(streams=True, storages=False):
            if len(entry) == 2 and entry[0] == "BodyText" and re.fullmatch(r"Section\d+", entry[1]):
                sections.append((int(entry[1][7:]), "/".join(entry)))
        if not sections:
            raise ExtractError(BROKEN_FILE, "HWP has no BodyText sections")

        left = MAX_HWP_BODY_BYTES
        for _, path in sorted(sections):
            if left <= 0:
                sink.truncated = True
                return
            raw = ole.openstream(path).read()
            if flags & HWP_COMPRESSED:
                inflater = zlib.decompressobj(-15)
                try:
                    body = inflater.decompress(raw, left)
                except zlib.error as e:
                    raise ExtractError(BROKEN_FILE, f"cannot inflate HWP section: {e}") from e
                cut = bool(inflater.unconsumed_tail) or len(body) >= left
            else:
                body, cut = raw[:left], len(raw) > left
            left -= len(body)
            if not hwp_section_paragraphs(body, sink):
                return
            if cut:
                sink.truncated = True
                return


def hwp_section_paragraphs(raw: bytes, sink: TextSink) -> bool:
    """레코드 헤더(tag 10bit · level 10bit · size 12bit, size 가 0xFFF 면 뒤 4바이트가 크기)를 따라가며 문단 글을 꺼낸다.

    sink 가 차서 멈췄으면 False.
    """
    pos = 0
    while pos + 4 <= len(raw):
        (head,) = struct.unpack_from("<I", raw, pos)
        pos += 4
        tag, size = head & 0x3FF, (head >> 20) & 0xFFF
        if size == 0xFFF:
            if pos + 4 > len(raw):
                break
            (size,) = struct.unpack_from("<I", raw, pos)
            pos += 4
        body = raw[pos:pos + size]
        pos += size
        if tag == HWPTAG_PARA_TEXT:
            text = hwp_para_text(body).strip()
            if text and not sink.add(text):
                return False
    return True


def hwp_para_text(body: bytes) -> str:
    """UTF-16LE 문단 글에서 제어 문자를 걸러 낸다.

    0·10·13·24~31 은 한 글자짜리 문자 제어(10·13 은 줄바꿈), 나머지 1~31 은 여덟 글자(16바이트)를 차지하는
    인라인·확장 제어(표·그림·각주 자리 등)라 건너뛴다. 9(탭)는 인라인 제어지만 칸을 띄우려고 탭으로 남긴다.
    """
    count = len(body) // 2
    codes = struct.unpack_from(f"<{count}H", body)
    chars, i = [], 0
    while i < count:
        code = codes[i]
        if code >= 32:
            chars.append(code)
            i += 1
        elif code in (10, 13):
            chars.append(10)
            i += 1
        elif code == 0 or 24 <= code <= 31:
            if code in (30, 31):
                chars.append(32)
            i += 1
        else:
            if code == 9:
                chars.append(9)
            i += 8
    return struct.pack(f"<{len(chars)}H", *chars).decode("utf-16-le", errors="replace")
